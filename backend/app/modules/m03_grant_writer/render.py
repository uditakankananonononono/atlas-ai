"""Deterministic proposal artifact rendering."""
from pathlib import Path
from html import escape

def render_docx_pdf(title:str,proposal:str,output_dir:Path)->dict[str,str]:
    from docx import Document
    from weasyprint import HTML
    output_dir.mkdir(parents=True,exist_ok=True)
    safe="proposal"
    docx_path=output_dir/f"{safe}.docx"; pdf_path=output_dir/f"{safe}.pdf"
    doc=Document(); doc.add_heading(title,0)
    for block in proposal.split("\n\n"):
        lines=block.strip().splitlines()
        if not lines: continue
        if len(lines)==1 and len(lines[0])<120 and (lines[0].isupper() or lines[0].endswith(":")): doc.add_heading(lines[0].rstrip(":"),level=1)
        else: doc.add_paragraph("\n".join(lines))
    doc.save(docx_path)
    paragraphs="".join(f"<p>{escape(p).replace(chr(10),'<br>')}</p>" for p in proposal.split("\n\n") if p.strip())
    HTML(string=f"<style>@page{{size:A4;margin:24mm}}body{{font:11pt Arial;line-height:1.45}}h1{{font-size:20pt}}</style><h1>{escape(title)}</h1>{paragraphs}").write_pdf(pdf_path)
    return {"docx":str(docx_path),"pdf":str(pdf_path)}
