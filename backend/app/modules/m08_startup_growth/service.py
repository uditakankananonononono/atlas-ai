"""Review-first generators for landing pages, decks, and code-grounded docs."""
from datetime import datetime,timezone
from hashlib import sha256
from io import BytesIO
import json,re,zipfile
from uuid import uuid4
from app.core.models import ApprovalRequest
from .schemas import *
MODULE_ID=8
class NotFoundError(LookupError):pass
class Service:
    def __init__(self,repository,approvals):self.repo=repository;self.approvals=approvals
    def _store(self,project_id,kind,files):
        stream=BytesIO()
        with zipfile.ZipFile(stream,"w",zipfile.ZIP_DEFLATED) as z:
            for name,content in sorted(files.items()):z.writestr(name,content)
        body=stream.getvalue(); ident=str(uuid4()); now=datetime.now(timezone.utc); manifest={"files":sorted(files),"bytes":len(body)}; digest=sha256(body).hexdigest()
        self.repo.save(id=ident,project_id=project_id,kind=kind,archive=body,manifest=manifest,sha256=digest,created_at=now)
        return BuildOut(id=ident,project_id=project_id,kind=kind,manifest=manifest,sha256=digest,created_at=now)
    def landing_page(self,data:LandingPageIn):
        features="\n".join(f"<li className='rounded-xl border p-5'>{x}</li>" for x in data.features)
        page=f'''export default function Page() {{ return <main className="mx-auto max-w-5xl p-8"><section className="py-24"><h1 className="text-6xl font-bold">{data.product_name}</h1><p className="mt-6 text-xl">{data.hero}</p></section><ul className="grid gap-6 md:grid-cols-3">{features}</ul><form action="/api/waitlist" method="post" className="mt-16 flex gap-2"><input required type="email" name="email" aria-label="Email" className="border p-3"/><button className="bg-black p-3 text-white">Join waitlist</button></form></main> }}'''
        route=f'''import {{createClient}} from "@supabase/supabase-js"; export async function POST(r:Request){{const f=await r.formData();const email=String(f.get("email")||"").trim().toLowerCase();if(!/^\\S+@\\S+\\.\\S+$/.test(email))return Response.json({{error:"invalid email"}},{{status:400}});const s=createClient(process.env.NEXT_PUBLIC_SUPABASE_URL!,process.env.SUPABASE_SERVICE_ROLE_KEY!);const {{error}}=await s.from("{data.waitlist_table}").upsert({{email}},{{onConflict:"email",ignoreDuplicates:true}});return error?Response.json({{error:"waitlist unavailable"}},{{status:503}}):Response.json({{ok:true}})}}'''
        files={"app/page.tsx":page,"app/api/waitlist/route.ts":route,"app/globals.css":"@tailwind base;\n@tailwind components;\n@tailwind utilities;\n",".env.example":"NEXT_PUBLIC_SUPABASE_URL=\nSUPABASE_SERVICE_ROLE_KEY=\n","README.md":"Generated review-first Next.js/Tailwind site. Configure Supabase, review, then request an approval-gated push/deploy.\n"}
        return self._store(data.project_id,"landing_page",files)
    def pitch_deck(self,data:PitchDeckIn):
        try:
            from pptx import Presentation
            from pptx.util import Inches
            prs=Presentation()
            for title,body in [(data.company,data.solution),("Problem",data.problem),("Solution",data.solution),("Market size",data.market_size),("Traction","\n".join(data.traction) or "No verified traction supplied"),("Ask",data.ask)]:
                slide=prs.slides.add_slide(prs.slide_layouts[1]);slide.shapes.title.text=title;slide.placeholders[1].text=body
            out=BytesIO();prs.save(out);files={"pitch-deck.pptx":out.getvalue(),"sources.json":json.dumps({"chart_series":data.chart_series,"note":"All claims supplied by user; verify before sharing."},indent=2)}
        except ImportError:
            files={"pitch-deck.md":f"# {data.company}\n## Problem\n{data.problem}\n## Solution\n{data.solution}\n## Market size\n{data.market_size}\n## Traction\n"+"\n".join(data.traction)+f"\n## Ask\n{data.ask}\n","sources.json":json.dumps(data.chart_series)}
        return self._store(data.project_id,"pitch_deck",files)
    def documentation(self,data:DocumentationIn):
        endpoints=[]
        for path,text in data.code_files.items():
            for method,route in re.findall(r'@(?:router|app)\.(get|post|put|patch|delete)\(["\']([^"\']+)',text):endpoints.append((method.upper(),route,path))
        paths={route:{method.lower():{"summary":f"Discovered in {path}","responses":{"200":{"description":"Success"}}}} for method,route,path in endpoints}
        openapi={"openapi":"3.1.0","info":{"title":data.title,"version":"0.1.0"},"paths":paths}
        files={"openapi.json":json.dumps(openapi,indent=2),"redoc.html":'<redoc spec-url="openapi.json"></redoc><script src="https://cdn.redoc.ly/redoc/latest/bundles/redoc.standalone.js"></script>',"USER_MANUAL.md":"# "+data.title+"\n\n"+"\n".join(f"- {x}" for x in data.feature_list)+"\n","TECHNICAL_BLOG.md":"# How "+data.title+" works\n\nThis article is grounded in the supplied code snapshot.\n\n"+"\n".join(f"- `{p}`" for p in sorted(data.code_files)),"evidence.json":json.dumps({"code_files":{p:sha256(c.encode()).hexdigest() for p,c in data.code_files.items()},"discovered_endpoints":endpoints},indent=2)}
        allowed={"openapi":{"openapi.json"},"redoc":{"redoc.html"},"user_manual":{"USER_MANUAL.md"},"technical_blog":{"TECHNICAL_BLOG.md"}};keep={"evidence.json"}|set().union(*(allowed[o] for o in data.outputs));return self._store(data.project_id,"documentation",{k:v for k,v in files.items() if k in keep})
    def propose(self,build_id,action):
        row=self.repo.get(build_id)
        if not row:raise NotFoundError(build_id)
        payload={"build_id":build_id,"project_id":row.project_id,"kind":row.kind,"sha256":row.sha256,"manifest":row.manifest};a=ApprovalRequest(id=str(uuid4()),module_id=MODULE_ID,action_type=action,payload=payload);self.approvals.put(a);return PublishProposal(approval_id=a.id,action_type=action,payload=payload)
