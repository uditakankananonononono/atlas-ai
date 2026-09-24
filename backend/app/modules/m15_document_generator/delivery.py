"""Approved document delivery: consume a render approval, render, verify, store, sign a download.

``routes.export`` files a ``render_document`` request with the Human Approval
Center, bound to ``tenant_id / document_id / version_id / format /
template_id / content_hash``. This module completes that flow:

1. the approval exists, belongs to the calling tenant, is a Module 15
   ``render_document`` item and is ``approved``;
2. the immutable version is loaded from the tenant's repository and its
   format, template and *recomputed* content hash must equal what was
   approved, so an edited or swapped version cannot ride on the approval;
3. Module 0 issues a one-shot permit bound to the exact reviewed payload
   (``Service.consume_effect``); a delivered approval cannot deliver twice;
4. the version is rendered with real engines - python-docx, python-pptx,
   Jinja2 LaTeX with every content value escaped, and ``pdflatex`` with
   shell-escape disabled and file access restricted for PDF - with
   matplotlib figures (alt text as caption) and a references list;
5. the output is checked to be a structurally valid file of that format,
   hashed, written to a tenant-scoped content-addressed store and a receipt
   is persisted;
6. downloads use an HMAC-signed, expiring token bound to tenant, approval
   and SHA-256; the bytes are re-hashed on every download.
"""
from __future__ import annotations

import base64
import hashlib
import hmac
import io
import json
import os
import re
import secrets
import shutil
import sqlite3
import subprocess
import tempfile
import threading
import time
import zipfile
from datetime import datetime, timezone
from pathlib import Path
from typing import Any
from uuid import uuid4

from .schemas import DocumentVersion, FigureSpec

ACTION = "render_document"
MODULE_ID = 15
TEMPLATES = Path(__file__).with_name("templates")
MIME = {"pdf": "application/pdf", "latex": "application/x-tex",
        "docx": "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
        "pptx": "application/vnd.openxmlformats-officedocument.presentationml.presentation"}
EXT = {"pdf": "pdf", "latex": "tex", "docx": "docx", "pptx": "pptx"}


class DeliveryError(RuntimeError):
    """Base error for approved delivery."""


class DeliveryNotFound(DeliveryError, KeyError):
    pass


class DeliveryForbidden(DeliveryError, PermissionError):
    pass


class DeliveryConflict(DeliveryError):
    pass


class RenderFailed(DeliveryError):
    pass


def _now() -> str:
    return datetime.now(timezone.utc).isoformat()


def _canonical(value: Any) -> str:
    return json.dumps(value, sort_keys=True, separators=(",", ":"), default=str)


def content_hash(content: dict[str, Any]) -> str:
    """Same hash Service.create_version stores."""
    return hashlib.sha256(json.dumps(content, sort_keys=True, default=str).encode()).hexdigest()


# --------------------------------------------------------------------------- rendering

_LATEX_ESCAPES = {"\\": r"\textbackslash{}", "&": r"\&", "%": r"\%", "$": r"\$", "#": r"\#",
                  "_": r"\_", "{": r"\{", "}": r"\}", "~": r"\textasciitilde{}", "^": r"\textasciicircum{}"}
_LATEX_RE = re.compile("|".join(re.escape(k) for k in _LATEX_ESCAPES))


def latex_escape(value: Any) -> str:
    return _LATEX_RE.sub(lambda m: _LATEX_ESCAPES[m.group()], str(value))


def _sections(content: dict[str, Any]) -> list[dict[str, str]]:
    sections = content.get("sections")
    if sections is None and content.get("body"):
        sections = [{"title": content.get("title", "Document"), "body": content["body"]}]
    out = []
    for item in sections or []:
        if not isinstance(item, dict) or not str(item.get("title", "")).strip():
            raise RenderFailed("every section needs a title")
        out.append({"title": str(item["title"]), "body": str(item.get("body", ""))})
    return out


def _reference(c: Any) -> str:
    parts = [", ".join(c.authors)] if c.authors else []
    parts.append(c.title)
    if c.url:
        parts.append(c.url)
    if c.accessed_at:
        parts.append(f"accessed {c.accessed_at}")
    return f"[{c.key}] " + ". ".join(parts)


def render_figure_png(spec: FigureSpec) -> bytes:
    """Render any figure spec with matplotlib (plotly specs are drawn with the same data)."""
    import matplotlib
    matplotlib.use("Agg")
    from matplotlib import pyplot as plt
    fig, ax = plt.subplots(figsize=(6, 3.6), dpi=120)
    try:
        if spec.kind == "pie":
            ax.pie(spec.data.get("values", []), labels=spec.data.get("labels", []))
        else:
            x, y = spec.data.get("x", []), spec.data.get("y", [])
            if len(x) != len(y):
                raise RenderFailed(f"figure {spec.id}: x and y lengths differ")
            {"line": ax.plot, "bar": ax.bar, "scatter": ax.scatter}[spec.kind](x, y)
        ax.set_title(str(spec.options.get("title", "")))
        out = io.BytesIO()
        fig.savefig(out, format="png", bbox_inches="tight")
        return out.getvalue()
    finally:
        plt.close(fig)


def _template(template_id: str, suffix: str) -> Path | None:
    if not re.fullmatch(r"[A-Za-z0-9_-]{1,100}", template_id):
        raise RenderFailed(f"invalid template id {template_id!r}")
    path = TEMPLATES / f"{template_id}{suffix}"
    return path if path.exists() else None


def render_latex_source(version: DocumentVersion, figure_files: list[dict[str, str]] | None = None) -> bytes:
    from jinja2 import Environment, FileSystemLoader, StrictUndefined
    template = _template(version.template_id, ".tex.j2")
    if template is None:
        raise RenderFailed(f"LaTeX template {version.template_id!r} does not exist")
    env = Environment(loader=FileSystemLoader(str(TEMPLATES)), undefined=StrictUndefined,
                      autoescape=False, finalize=lambda v: v)
    content = version.content
    context = {
        "title": latex_escape(content.get("title", version.document_id)),
        "author": latex_escape(content.get("author", "Atlas AI")),
        "sections": [{"title": latex_escape(s["title"]), "body": latex_escape(s["body"])} for s in _sections(content)],
        "figures": [{"file": f["file"], "caption": latex_escape(f["caption"])} for f in figure_files or []],
        "references": [latex_escape(_reference(c)) for c in version.citations],
    }
    return env.get_template(template.name).render(**context).encode()


def render_pdf(version: DocumentVersion, timeout: int = 90) -> bytes:
    exe = shutil.which("pdflatex")
    if not exe:
        raise RenderFailed("pdflatex is not installed")
    work = Path(tempfile.mkdtemp(prefix="atlas-m15-pdf-"))
    try:
        figures = []
        for index, spec in enumerate(version.figures):
            name = f"figure{index}.png"
            (work / name).write_bytes(render_figure_png(spec))
            figures.append({"file": name, "caption": spec.alt_text})
        (work / "doc.tex").write_bytes(render_latex_source(version, figures))
        env = {"PATH": "/usr/bin:/bin", "HOME": str(work), "openin_any": "p", "openout_any": "p",
               "shell_escape": "f", "TEXMFOUTPUT": str(work)}
        if os.getenv("ATLAS_TEXMFHOME"):  # optional extra read-only package tree
            env["TEXMFHOME"] = os.environ["ATLAS_TEXMFHOME"]
        proc = subprocess.run([exe, "-no-shell-escape", "-halt-on-error", "-interaction=nonstopmode", "doc.tex"],
                              cwd=work, env=env, capture_output=True, timeout=timeout)
        pdf = work / "doc.pdf"
        if proc.returncode != 0 or not pdf.exists():
            tail = proc.stdout.decode(errors="replace")[-800:]
            raise RenderFailed(f"pdflatex failed: {tail}")
        return pdf.read_bytes()
    except subprocess.TimeoutExpired as exc:
        raise RenderFailed("pdflatex timed out") from exc
    finally:
        shutil.rmtree(work, ignore_errors=True)


def render_docx(version: DocumentVersion) -> bytes:
    from docx import Document
    from docx.shared import Inches
    template = _template(version.template_id, ".docx")
    doc = Document(str(template)) if template else Document()
    content = version.content
    doc.core_properties.title = str(content.get("title", version.document_id))
    doc.core_properties.author = str(content.get("author", "Atlas AI"))
    doc.add_heading(str(content.get("title", version.document_id)), level=0)
    for section in _sections(content):
        doc.add_heading(section["title"], level=1)
        for para in section["body"].split("\n\n"):
            if para.strip():
                doc.add_paragraph(para.strip())
    for spec in version.figures:
        doc.add_picture(io.BytesIO(render_figure_png(spec)), width=Inches(5.5))
        doc.add_paragraph(spec.alt_text).style = doc.styles["Caption"]
    if version.citations:
        doc.add_heading("References", level=1)
        for c in version.citations:
            doc.add_paragraph(_reference(c), style="List Number")
    out = io.BytesIO()
    doc.save(out)
    return out.getvalue()


def render_pptx(version: DocumentVersion) -> bytes:
    from pptx import Presentation
    from pptx.util import Inches
    template = _template(version.template_id, ".pptx")
    prs = Presentation(str(template)) if template else Presentation()
    slides = version.content.get("slides") or []
    if not slides:
        raise RenderFailed("PPTX has no slides")
    if version.content.get("title"):
        cover = prs.slides.add_slide(prs.slide_layouts[0])
        cover.shapes.title.text = str(version.content["title"])
        if len(cover.placeholders) > 1:
            cover.placeholders[1].text = str(version.content.get("author", ""))
    for item in slides:
        if not item.get("title"):
            raise RenderFailed("every slide requires a title")
        slide = prs.slides.add_slide(prs.slide_layouts[1])
        slide.shapes.title.text = str(item["title"])
        lines = [str(b) for b in item.get("bullets", [])] or [p for p in str(item.get("body", "")).split("\n") if p]
        slide.placeholders[1].text = "\n".join(lines)
        if item.get("notes"):
            slide.notes_slide.notes_text_frame.text = str(item["notes"])
    for spec in version.figures:
        slide = prs.slides.add_slide(prs.slide_layouts[5])
        slide.shapes.title.text = str(spec.options.get("title") or spec.id)
        pic = slide.shapes.add_picture(io.BytesIO(render_figure_png(spec)), Inches(1), Inches(1.6), width=Inches(8))
        pic._element.nvPicPr.cNvPr.set("descr", spec.alt_text)
    if version.citations:
        refs = prs.slides.add_slide(prs.slide_layouts[1])
        refs.shapes.title.text = "References"
        refs.placeholders[1].text = "\n".join(_reference(c) for c in version.citations)
    out = io.BytesIO()
    prs.save(out)
    return out.getvalue()


def render_version(version: DocumentVersion) -> bytes:
    if version.format == "latex":
        return render_latex_source(version)
    return {"pdf": render_pdf, "docx": render_docx, "pptx": render_pptx}[version.format](version)


def validate_output(fmt: str, data: bytes) -> dict[str, Any]:
    """Structural check that the bytes are really a file of the approved format."""
    if not data:
        raise RenderFailed("renderer returned no bytes")
    if fmt == "pdf":
        if not data.startswith(b"%PDF-") or b"%%EOF" not in data[-2048:]:
            raise RenderFailed("output is not a complete PDF")
        try:
            import fitz  # PyMuPDF, a declared project dependency
        except ImportError:
            return {"kind": "pdf", "pages": None, "parsed": False}
        try:
            with fitz.open(stream=data, filetype="pdf") as pdf:
                pages = pdf.page_count
                images = sum(len(page.get_images()) for page in pdf)
        except Exception as exc:
            raise RenderFailed(f"output PDF does not parse: {exc}") from exc
        if pages < 1:
            raise RenderFailed("output PDF has no pages")
        return {"kind": "pdf", "pages": pages, "images": images, "parsed": True}
    if fmt == "latex":
        text = data.decode("utf-8")
        if "\\begin{document}" not in text or "\\end{document}" not in text:
            raise RenderFailed("output is not a complete LaTeX document")
        return {"kind": "latex", "characters": len(text)}
    try:
        names = zipfile.ZipFile(io.BytesIO(data)).namelist()
    except zipfile.BadZipFile as exc:
        raise RenderFailed(f"output is not a valid {fmt} package") from exc
    main = "word/document.xml" if fmt == "docx" else "ppt/presentation.xml"
    if main not in names:
        raise RenderFailed(f"output is missing {main}")
    info: dict[str, Any] = {"kind": fmt, "parts": len(names)}
    if fmt == "pptx":
        info["slides"] = len([n for n in names if re.fullmatch(r"ppt/slides/slide\d+\.xml", n)])
    return info


# --------------------------------------------------------------------------- storage + signing

class DeliveryStore:
    def __init__(self, root: str | Path, signing_key: bytes | None = None) -> None:
        self.root = Path(root)
        self.root.mkdir(parents=True, exist_ok=True)
        self.db = self.root / "m15-deliveries.sqlite"
        self._lock = threading.Lock()
        with sqlite3.connect(self.db) as db:
            db.execute("CREATE TABLE IF NOT EXISTS m15_deliveries (tenant_id TEXT NOT NULL, approval_id TEXT NOT NULL, "
                       "receipt TEXT NOT NULL, created_at TEXT NOT NULL, PRIMARY KEY (tenant_id, approval_id))")
        self.signing_key = signing_key or self._load_key()

    def _load_key(self) -> bytes:
        env = os.getenv("ATLAS_DOWNLOAD_SIGNING_KEY")
        if env:
            return env.encode()
        path = self.root / "download-signing.key"
        if not path.exists():
            fd = os.open(path, os.O_WRONLY | os.O_CREAT | os.O_EXCL, 0o600)
            with os.fdopen(fd, "wb") as handle:
                handle.write(secrets.token_bytes(32))
        return path.read_bytes()

    def _blob(self, tenant_id: str, sha: str) -> Path:
        return self.root / "files" / hashlib.sha256(tenant_id.encode()).hexdigest()[:24] / sha[:2] / sha

    def put(self, tenant_id: str, data: bytes) -> str:
        sha = hashlib.sha256(data).hexdigest()
        path = self._blob(tenant_id, sha)
        if not path.exists():
            path.parent.mkdir(parents=True, exist_ok=True)
            tmp = path.with_suffix(".tmp")
            tmp.write_bytes(data)
            tmp.replace(path)
        return sha

    def read(self, tenant_id: str, sha: str) -> bytes:
        path = self._blob(tenant_id, sha)
        if not path.exists():
            raise DeliveryNotFound(sha)
        data = path.read_bytes()
        if hashlib.sha256(data).hexdigest() != sha:
            raise DeliveryError("stored file failed its integrity check")
        return data

    def save(self, tenant_id: str, receipt: dict[str, Any]) -> None:
        with self._lock, sqlite3.connect(self.db) as db:
            try:
                db.execute("INSERT INTO m15_deliveries VALUES (?,?,?,?)",
                           (tenant_id, receipt["approval_id"], _canonical(receipt), _now()))
            except sqlite3.IntegrityError as exc:
                raise DeliveryConflict("approval was already delivered") from exc

    def get(self, tenant_id: str, approval_id: str) -> dict[str, Any] | None:
        with sqlite3.connect(self.db) as db:
            row = db.execute("SELECT receipt FROM m15_deliveries WHERE tenant_id=? AND approval_id=?",
                             (tenant_id, approval_id)).fetchone()
        return json.loads(row[0]) if row else None

    def list(self, tenant_id: str, limit: int = 100) -> list[dict[str, Any]]:
        with sqlite3.connect(self.db) as db:
            rows = db.execute("SELECT receipt FROM m15_deliveries WHERE tenant_id=? ORDER BY created_at DESC LIMIT ?",
                              (tenant_id, limit)).fetchall()
        return [json.loads(r[0]) for r in rows]

    # tokens: base64url(json).base64url(hmac-sha256)
    def sign(self, tenant_id: str, approval_id: str, sha: str, ttl_seconds: int) -> tuple[str, int]:
        expires = int(time.time()) + ttl_seconds
        body = base64.urlsafe_b64encode(_canonical({"t": tenant_id, "a": approval_id, "s": sha, "e": expires}).encode()).rstrip(b"=")
        mac = base64.urlsafe_b64encode(hmac.new(self.signing_key, body, hashlib.sha256).digest()).rstrip(b"=")
        return f"{body.decode()}.{mac.decode()}", expires

    def verify(self, token: str) -> dict[str, Any]:
        try:
            body, mac = token.split(".", 1)
            expected = base64.urlsafe_b64encode(hmac.new(self.signing_key, body.encode(), hashlib.sha256).digest()).rstrip(b"=").decode()
            if not hmac.compare_digest(mac, expected):
                raise DeliveryForbidden("invalid download signature")
            claims = json.loads(base64.urlsafe_b64decode(body + "=" * (-len(body) % 4)))
        except (ValueError, json.JSONDecodeError) as exc:
            raise DeliveryForbidden("malformed download token") from exc
        if int(claims["e"]) < time.time():
            raise DeliveryForbidden("download link has expired")
        return claims


# --------------------------------------------------------------------------- service

class ApprovedDeliveryService:
    def __init__(self, tenant_id: str, *, actor_id: str | None = None, center: Any = None,
                 versions: Any = None, store: DeliveryStore | None = None,
                 link_ttl_seconds: int | None = None, download_prefix: str = "/api/v1/document-generator/downloads") -> None:
        if center is None:
            from app.modules.m00_approval_center.service import default_service
            center = default_service()
        if versions is None:
            from .sql_repository import SqlVersionRepository
            versions = SqlVersionRepository(tenant_id)
        if store is None:
            store = DeliveryStore(Path(os.getenv("ATLAS_RUNTIME_DATA_DIR", "/tmp/atlas-runtime")) / "m15-delivery")
        self.tenant_id = tenant_id
        self.actor_id = actor_id or tenant_id
        self.center, self.versions, self.store = center, versions, store
        self.link_ttl = link_ttl_seconds or int(os.getenv("ATLAS_DOWNLOAD_LINK_TTL_SECONDS", "3600"))
        self.download_prefix = download_prefix

    def _link(self, receipt: dict[str, Any]) -> dict[str, Any]:
        token, expires = self.store.sign(self.tenant_id, receipt["approval_id"], receipt["sha256"], self.link_ttl)
        return {**receipt, "download_url": f"{self.download_prefix}/{token}",
                "download_expires_at": datetime.fromtimestamp(expires, timezone.utc).isoformat()}

    def deliver(self, approval_id: str) -> dict[str, Any]:
        try:
            view = self.center.get(approval_id)
        except KeyError as exc:
            raise DeliveryNotFound(approval_id) from exc
        if view.get("user_id") != self.tenant_id:
            raise DeliveryNotFound(approval_id)
        if view.get("module_id") != MODULE_ID or view.get("action_type") != ACTION:
            raise DeliveryForbidden("approval is not a Module 15 render_document item")
        status = getattr(view.get("status"), "value", view.get("status"))
        if status != "approved":
            raise DeliveryForbidden(f"approval is {status}, not approved")
        if self.store.get(self.tenant_id, approval_id):
            raise DeliveryConflict("approval was already delivered; read it back for a fresh link")
        payload = dict(view.get("payload") or {})
        if payload.get("tenant_id") != self.tenant_id:
            raise DeliveryForbidden("approval payload is bound to another tenant")
        version = self.versions.get(str(payload.get("version_id")))
        if version is None:
            raise DeliveryNotFound("approved document version no longer exists")
        recomputed = content_hash(version.content)
        mismatches = [k for k, have in (("document_id", version.document_id), ("format", version.format),
                                         ("template_id", version.template_id), ("content_hash", recomputed))
                      if payload.get(k) != have]
        if mismatches:
            raise DeliveryForbidden(f"version does not match the approved request: {', '.join(mismatches)}")
        consumed = any(e.get("event") == "effect_consumed" for e in self.center.audit(approval_id))
        if consumed:
            raise DeliveryConflict("approval permit was already consumed")
        # Render before consuming: a render failure must not burn the human's approval.
        started = time.monotonic()
        data = render_version(version)
        check = validate_output(version.format, data)
        try:
            permit = self.center.consume_effect(approval_id, module_id=MODULE_ID, action_type=ACTION,
                                                payload=payload, user_id=self.tenant_id,
                                                effect_id=f"m15-deliver:{approval_id}", actor=self.actor_id)
        except Exception as exc:
            if type(exc).__name__ == "ApprovalConflictError":
                raise DeliveryConflict(str(exc)) from exc
            raise
        sha = self.store.put(self.tenant_id, data)
        filename = f"{re.sub(r'[^A-Za-z0-9._-]', '_', version.document_id)[:80]}-v{version.version_number}.{EXT[version.format]}"
        receipt = {"delivery_id": str(uuid4()), "approval_id": approval_id, "tenant_id": self.tenant_id,
                   "actor_id": self.actor_id, "document_id": version.document_id, "version_id": version.id,
                   "version_number": version.version_number, "format": version.format,
                   "template_id": version.template_id, "approved_content_hash": payload["content_hash"],
                   "sha256": sha, "byte_size": len(data), "mime_type": MIME[version.format],
                   "filename": filename, "validation": check, "access": "private",
                   "render_seconds": round(time.monotonic() - started, 3),
                   "permit_consumed_at": str(permit.get("consumed_at")), "delivered_at": _now()}
        receipt["receipt_sha256"] = hashlib.sha256(_canonical(receipt).encode()).hexdigest()
        self.store.save(self.tenant_id, receipt)
        return self._link(receipt)

    def readback(self, approval_id: str) -> dict[str, Any]:
        receipt = self.store.get(self.tenant_id, approval_id)
        if receipt is None:
            raise DeliveryNotFound(approval_id)
        return self._link(receipt)

    def list(self, limit: int = 100) -> list[dict[str, Any]]:
        return self.store.list(self.tenant_id, limit)

    def download(self, token: str) -> tuple[bytes, dict[str, Any]]:
        claims = self.store.verify(token)
        if claims["t"] != self.tenant_id:
            raise DeliveryNotFound("download")
        receipt = self.store.get(self.tenant_id, claims["a"])
        if receipt is None or receipt["sha256"] != claims["s"]:
            raise DeliveryNotFound("download")
        return self.store.read(self.tenant_id, claims["s"]), receipt


def verify_receipt(receipt: dict[str, Any]) -> bool:
    keys = {"download_url", "download_expires_at", "receipt_sha256"}
    body = {k: v for k, v in receipt.items() if k not in keys}
    return receipt.get("receipt_sha256") == hashlib.sha256(_canonical(body).encode()).hexdigest()
