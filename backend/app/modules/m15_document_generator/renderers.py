"""Optional native render adapters. Imports are lazy so the API works without render extras."""
from io import BytesIO
from pathlib import Path

def render_docx(template:Path,context:dict)->bytes:
    from docxtpl import DocxTemplate
    document=DocxTemplate(str(template));document.render(context);out=BytesIO();document.save(out);return out.getvalue()
def render_pptx(template:Path,content:dict)->bytes:
    from pptx import Presentation
    prs=Presentation(str(template))
    for item in content.get("slides",[]):
        slide=prs.slides.add_slide(prs.slide_layouts[1]);slide.shapes.title.text=item["title"];slide.placeholders[1].text="\n".join(item.get("bullets",[]))
    out=BytesIO();prs.save(out);return out.getvalue()
