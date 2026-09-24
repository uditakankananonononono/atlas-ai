"""Module 15 approved delivery: real renders, approval binding, one-shot permit, signed downloads."""
from __future__ import annotations

import hashlib
import io
import shutil
import zipfile

import pytest
from fastapi import FastAPI
from fastapi.testclient import TestClient
from sqlalchemy import create_engine
from sqlalchemy.orm import sessionmaker

from app.core.database import Base
from app.core.models import ApprovalStatus
from app.modules.m00_approval_center.service import ApprovalBroadcaster, Service as Center
from app.modules.m15_document_generator import routes
from app.modules.m15_document_generator.delivery import (
    ApprovedDeliveryService, DeliveryConflict, DeliveryForbidden, DeliveryNotFound, DeliveryStore,
    RenderFailed, latex_escape, render_latex_source, validate_output, verify_receipt,
)
from app.modules.m15_document_generator.schemas import CreateVersionRequest
from app.modules.m15_document_generator.service import Service as Documents
from app.modules.m15_document_generator.sql_repository import SqlVersionRepository


@pytest.fixture
def env(tmp_path, monkeypatch):
    engine = create_engine(f"sqlite:///{tmp_path}/db.sqlite", connect_args={"check_same_thread": False})
    Base.metadata.create_all(engine)
    factory = sessionmaker(bind=engine, expire_on_commit=False)
    center = Center(session_factory=factory, broadcaster=ApprovalBroadcaster())
    import app.core.approvals as facade
    monkeypatch.setattr(facade, "default_service", lambda: center)

    class Env:
        pass
    e = Env()
    e.center, e.factory, e.tmp = center, factory, tmp_path
    e.repo = lambda tenant="tenant-a": SqlVersionRepository(tenant, session_factory=factory)
    e.docs = lambda tenant="tenant-a": Documents(facade.approvals, repository=e.repo(tenant))
    e.store = DeliveryStore(tmp_path / "store", signing_key=b"k" * 32)
    e.svc = lambda tenant="tenant-a", **kw: ApprovedDeliveryService(tenant, center=center, versions=e.repo(tenant), store=e.store, **kw)
    return e


REPORT = {"title": "Tahoe clarity 50% & rising", "author": "Udita",
          "sections": [{"title": "Method", "body": "Secchi depth_m ~ 21 {weekly} $0 cost #1"},
                       {"title": "Result", "body": "Clarity improved.\n\nSecond paragraph."}]}
FIG = {"id": "f1", "engine": "matplotlib", "kind": "line", "data": {"x": [1, 2, 3], "y": [20, 21, 22]},
       "options": {"title": "Depth"}, "alt_text": "Secchi depth rising over three weeks"}
CITE = {"key": "trg2025", "title": "State of the Lake", "url": "https://tahoe.ucdavis.edu/stateofthelake", "authors": ["UC Davis TERC"]}


def approve(env, fmt, content, tenant="tenant-a", template="report", approve=True, figures=(), citations=()):
    docs = env.docs(tenant)
    version = docs.create_version(tenant, "tahoe-report", CreateVersionRequest(
        title="t", format=fmt, template_id=template, content=content, figures=list(figures), citations=list(citations)))
    proposal = docs.propose_export(version)
    if approve:
        env.center.decide(proposal.approval_id, ApprovalStatus.APPROVED, decided_by="udita")
    return proposal.approval_id, version


def _pdflatex_works() -> bool:
    """pdflatex present AND the report template's packages (geometry, graphicx, hyperref) resolve."""
    if not shutil.which("pdflatex"):
        return False
    import subprocess
    import os
    env = {**os.environ, **({"TEXMFHOME": os.environ["ATLAS_TEXMFHOME"]} if os.getenv("ATLAS_TEXMFHOME") else {})}
    probe = subprocess.run(["kpsewhich", "geometry.sty", "graphicx.sty", "hyperref.sty", "pdftexcmds.sty"],
                           capture_output=True, text=True, env=env)
    return len(probe.stdout.split()) == 4


needs_pdflatex = pytest.mark.skipif(not _pdflatex_works(), reason="pdflatex or required LaTeX packages missing")


@needs_pdflatex
def test_pdf_delivery_renders_escaped_content_figures_and_references(env):
    approval_id, version = approve(env, "pdf", REPORT, figures=[FIG], citations=[CITE])
    svc = env.svc()
    receipt = svc.deliver(approval_id)
    assert receipt["format"] == "pdf" and receipt["mime_type"] == "application/pdf"
    assert receipt["validation"]["kind"] == "pdf" and receipt["validation"]["pages"] >= 1
    assert receipt["approved_content_hash"] == version.content_hash and verify_receipt(receipt)
    data, _ = svc.download(receipt["download_url"].rsplit("/", 1)[1])
    assert hashlib.sha256(data).hexdigest() == receipt["sha256"] and data.startswith(b"%PDF-")
    assert receipt["validation"]["parsed"] and receipt["validation"]["images"] == 1
    import fitz
    with fitz.open(stream=data, filetype="pdf") as pdf:
        text = "".join(page.get_text() for page in pdf)
    assert "Tahoe clarity 50% & rising" in text and "Secchi depth_m ~ 21 {weekly} $0 cost #1" in text
    assert "Secchi depth rising over three weeks" in text and "[trg2025] UC Davis TERC. State of the Lake" in text


def test_latex_escaping_blocks_command_injection(env):
    hostile = {"title": r"\input{/etc/passwd}", "sections": [{"title": "x", "body": r"\write18{rm -rf /} & $x$"}]}
    _, version = approve(env, "latex", hostile, approve=False)
    tex = render_latex_source(version).decode()
    assert r"\input{/etc/passwd}" not in tex and r"\textbackslash{}input\{/etc/passwd\}" in tex
    assert r"\write18" not in tex
    assert latex_escape("50% & $") == r"50\% \& \$"


def test_docx_and_pptx_are_real_packages_with_content(env):
    svc = env.svc()
    a1, _ = approve(env, "docx", REPORT, figures=[FIG], citations=[CITE])
    docx_receipt = svc.deliver(a1)
    data, _ = svc.download(docx_receipt["download_url"].rsplit("/", 1)[1])
    from docx import Document
    doc = Document(io.BytesIO(data))
    text = "\n".join(p.text for p in doc.paragraphs)
    assert "Secchi depth_m ~ 21 {weekly} $0 cost #1" in text and "Second paragraph." in text
    assert "[trg2025] UC Davis TERC. State of the Lake" in text
    assert any(n.startswith("word/media/") for n in zipfile.ZipFile(io.BytesIO(data)).namelist())

    deck = {"title": "Tahoe", "slides": [{"title": "Why", "bullets": ["clarity", "algae"], "notes": "speak slowly"},
                                         {"title": "How", "body": "line one\nline two"}]}
    a2, _ = approve(env, "pptx", deck, figures=[FIG])
    pptx_receipt = svc.deliver(a2)
    data, _ = svc.download(pptx_receipt["download_url"].rsplit("/", 1)[1])
    from pptx import Presentation
    prs = Presentation(io.BytesIO(data))
    titles = [s.shapes.title.text for s in prs.slides if s.shapes.title is not None]
    assert titles[:3] == ["Tahoe", "Why", "How"] and pptx_receipt["validation"]["slides"] == 4
    assert prs.slides[1].notes_slide.notes_text_frame.text == "speak slowly"


def test_guards_pending_denied_foreign_and_wrong_action(env):
    svc = env.svc()
    pending, _ = approve(env, "latex", REPORT, approve=False)
    with pytest.raises(DeliveryForbidden, match="pending"):
        svc.deliver(pending)
    denied, _ = approve(env, "latex", REPORT, approve=False)
    env.center.decide(denied, ApprovalStatus.DENIED, decided_by="udita")
    with pytest.raises(DeliveryForbidden, match="denied"):
        svc.deliver(denied)
    foreign, _ = approve(env, "latex", REPORT, tenant="tenant-b")
    with pytest.raises(DeliveryNotFound):
        svc.deliver(foreign)
    other = env.center.submit(module_id=15, action_type="publish", payload={"tenant_id": "tenant-a"}, user_id="tenant-a")
    env.center.decide(other["id"], ApprovalStatus.APPROVED, decided_by="udita")
    with pytest.raises(DeliveryForbidden, match="not a Module 15"):
        svc.deliver(other["id"])


def test_approval_bound_to_content_hash_rejects_swapped_version(env):
    svc = env.svc()
    view = env.center.submit(module_id=15, action_type="render_document", user_id="tenant-a", payload={
        "tenant_id": "tenant-a", "document_id": "tahoe-report", "version_id": "missing", "format": "latex",
        "template_id": "report", "content_hash": "0" * 64})
    env.center.decide(view["id"], ApprovalStatus.APPROVED, decided_by="udita")
    with pytest.raises(DeliveryNotFound):
        svc.deliver(view["id"])
    _, version = approve(env, "latex", REPORT, approve=False)
    forged = env.center.submit(module_id=15, action_type="render_document", user_id="tenant-a", payload={
        "tenant_id": "tenant-a", "document_id": "tahoe-report", "version_id": version.id, "format": "latex",
        "template_id": "report", "content_hash": "f" * 64})
    env.center.decide(forged["id"], ApprovalStatus.APPROVED, decided_by="udita")
    with pytest.raises(DeliveryForbidden, match="content_hash"):
        svc.deliver(forged["id"])


def test_one_shot_delivery_and_render_failure_does_not_burn_approval(env):
    svc = env.svc()
    ok, _ = approve(env, "latex", REPORT)
    first = svc.deliver(ok)
    with pytest.raises(DeliveryConflict):
        svc.deliver(ok)
    other_store = ApprovedDeliveryService("tenant-a", center=env.center, versions=env.repo(),
                                          store=DeliveryStore(env.tmp / "store2", signing_key=b"k" * 32))
    with pytest.raises(DeliveryConflict, match="consumed"):
        other_store.deliver(ok)
    again = svc.readback(ok)
    assert again["sha256"] == first["sha256"] and again["download_url"] != "" and verify_receipt(again)
    bad, _ = approve(env, "latex", REPORT, template="missing_template")
    with pytest.raises(RenderFailed):
        svc.deliver(bad)
    assert "effect_consumed" not in [e["event"] for e in env.center.audit(bad)]


def test_download_tokens_are_signed_expiring_and_tenant_bound(env):
    svc = env.svc(link_ttl_seconds=60)
    approval_id, _ = approve(env, "latex", REPORT)
    token = svc.deliver(approval_id)["download_url"].rsplit("/", 1)[1]
    body, mac = token.split(".")
    with pytest.raises(DeliveryForbidden, match="signature"):
        svc.download(body + "." + mac[:-2] + ("AA" if not mac.endswith("AA") else "BB"))
    with pytest.raises(DeliveryNotFound):
        env.svc("tenant-b").download(token)
    expired, _ = env.store.sign("tenant-a", approval_id, "0" * 64, -5)
    with pytest.raises(DeliveryForbidden, match="expired"):
        svc.download(expired)


def test_output_validation_rejects_fake_files():
    with pytest.raises(RenderFailed):
        validate_output("pdf", b"not a pdf")
    with pytest.raises(RenderFailed):
        validate_output("docx", b"PK not really")
    with pytest.raises(RenderFailed):
        validate_output("latex", b"hello")


def test_http_flow_export_approve_deliver_download(env):
    app = FastAPI(); app.include_router(routes.router)
    app.dependency_overrides[routes.get_service] = lambda: env.docs("tenant-a")
    app.dependency_overrides[routes.get_delivery_service] = lambda: env.svc("tenant-a")
    client = TestClient(app); h = {"X-Atlas-Tenant": "tenant-a"}
    version = client.post("/document-generator/documents/tahoe/versions", headers=h, json={
        "title": "t", "format": "docx", "template_id": "report", "content": REPORT}).json()
    approval_id = client.post(f"/document-generator/versions/{version['id']}/export-proposals", headers=h).json()["approval_id"]
    assert client.post(f"/document-generator/approvals/{approval_id}/deliver", headers=h).status_code == 403
    env.center.decide(approval_id, ApprovalStatus.APPROVED, decided_by="udita")
    delivered = client.post(f"/document-generator/approvals/{approval_id}/deliver", headers=h)
    assert delivered.status_code == 201, delivered.text
    url = delivered.json()["download_url"].replace("/api/v1", "")
    file = client.get(url, headers=h)
    assert file.status_code == 200 and hashlib.sha256(file.content).hexdigest() == delivered.json()["sha256"]
    assert file.headers["content-disposition"].endswith('tahoe-v1.docx"')
    assert client.post(f"/document-generator/approvals/{approval_id}/deliver", headers=h).status_code == 409
    assert client.get(f"/document-generator/deliveries/{approval_id}", headers=h).status_code == 200
    assert [d["approval_id"] for d in client.get("/document-generator/deliveries", headers=h).json()] == [approval_id]
    assert client.get("/document-generator/downloads/garbage", headers=h).status_code == 403
